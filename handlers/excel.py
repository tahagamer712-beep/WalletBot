"""
تصدير العمليات إلى ملف Excel (openpyxl) — تقرير احترافي بورقة واحدة (Sheet
وحدة)، مقسّم لأقسام واضحة بعناوين ملوّنة: الملخص، المعاملات، التصنيفات
(مصروف)، الدخل حسب التصنيف، الالتزامات، الدفعات والاستقطاعات، أهداف الادخار.
"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from telegram import Update, InlineKeyboardMarkup, InlineKeyboardButton
from telegram.ext import ContextTypes

import config
import database.db as db
from reports.balance import balance_before
from utils import dates, categories as cat_utils
from utils.reply import nav_render, back_row
from reports.statistics import period_summary

ARABIC_MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
                 "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]

# هوية بصرية بسيطة تتماشى مع لون البوت البنفسجي بدل شكل "جدول افتراضي رمادي"
COLOR_HEADER = "6C3483"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_ALT_ROW = "F4ECF7"
COLOR_EXPENSE = "C0392B"
COLOR_INCOME = "1E8449"
COLOR_TOTAL_FILL = "D7BDE2"
COLOR_SECTION = "9B59B6"      # بنفسجي أفتح لعناوين الأقسام (تمييز عن عنوان التقرير الرئيسي)

MAIN_COLS = 8  # عرض الجدول الأوسع (جدول المعاملات) — كل الأقسام تُحاذى عليه
THIN_BORDER = Border(*(Side(style="thin", color="CCCCCC"),) * 4)
HEADER_FONT = Font(name="Calibri", bold=True, color=COLOR_HEADER_TEXT, size=11)
HEADER_FILL = PatternFill("solid", fgColor=COLOR_HEADER)
SECTION_FONT = Font(name="Calibri", bold=True, color=COLOR_HEADER_TEXT, size=12)
SECTION_FILL = PatternFill("solid", fgColor=COLOR_SECTION)
TITLE_FONT = Font(name="Calibri", bold=True, size=15, color=COLOR_HEADER)


def _sanitize_label(label: str) -> str:
    cleaned = re.sub(r"[^\w\u0600-\u06FF]+", "_", label.strip())
    return cleaned[:40] or "مستخدم"


async def show_excel_menu(update, context: ContextTypes.DEFAULT_TYPE, user: dict):
    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("📁 الشهر الحالي", callback_data="xls:month:0")],
        [InlineKeyboardButton("📁 الشهر الماضي", callback_data="xls:month:1")],
        [InlineKeyboardButton("📁 كل الأشهر", callback_data="xls:all")],
        back_row("main:home"),
    ])
    await nav_render(update, context, "اختر الفترة المطلوب تصديرها:", reply_markup=kb)


# ------------------------------------ أدوات تنسيق مشتركة ------------------------------------

def _section_header(ws, row, text, span=MAIN_COLS):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 22
    return row + 1


def _table_header(ws, row, headers):
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    return row + 1


def _data_row(ws, row, values, alt=False, bold_cols=(), fills=None):
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if alt:
            cell.fill = PatternFill("solid", fgColor=COLOR_ALT_ROW)
        if col_idx in bold_cols:
            cell.font = Font(bold=True)
        if fills and col_idx in fills:
            cell.fill = PatternFill("solid", fgColor=fills[col_idx])
    return row + 1


def _blank(row, n=1):
    return row + n


# ------------------------------------ قسم 1: الملخص ------------------------------------

def _write_summary_section(ws, row, summary, total_deductions, total_saved_period,
                           opening_iqd=0.0, closing_iqd=None):
    row = _section_header(ws, row, "📊 الملخص")
    row = _table_header(ws, row, ["البند", "القيمة (IQD)"])
    net_month = summary["net_iqd"] - total_deductions
    if closing_iqd is None:
        closing_iqd = opening_iqd + net_month
    items = [
        ("🔄 الرصيد المرحّل (بداية الفترة)", opening_iqd, False),
        ("💵 إجمالي الدخل", summary["income_iqd"], False),
        ("💸 إجمالي المصروفات", summary["expense_iqd"], False),
        ("📌 إجمالي الاستقطاعات (سلف/أقساط/التزامات)", total_deductions, False),
        ("🎯 إجمالي الادخار بالفترة", total_saved_period, False),
        ("📌 صافي الفترة", net_month, False),
        ("✅ رصيد نهاية الفترة", closing_iqd, True),
    ]
    for i, (label, value, is_closing) in enumerate(items):
        fills = {1: COLOR_TOTAL_FILL, 2: COLOR_TOTAL_FILL} if is_closing else None
        row = _data_row(ws, row, [label, value], alt=(i % 2 == 1),
                        bold_cols=(1, 2) if is_closing else (), fills=fills)
        ws.cell(row=row - 1, column=2).number_format = "#,##0"
    return _blank(row)


# ------------------------------------ قسم 2: المعاملات ------------------------------------

def _write_transactions_section(ws, row, rows):
    row = _section_header(ws, row, "📋 جميع المعاملات")
    headers = ["التاريخ", "الوقت", "النوع", "التصنيف", "المبلغ", "العملة", "بالدينار", "الملاحظة"]
    row = _table_header(ws, row, headers)
    for i, tx in enumerate(rows):
        type_label = "مصروف" if tx["type"] == "expense" else "دخل"
        icon = cat_utils.icon(tx["category"])
        values = [tx["transaction_date"], tx["transaction_time"], type_label, f"{icon} {tx['category']}",
                  tx["amount"], tx["currency"], tx["amount_iqd"], tx["note"] or ""]
        row = _data_row(ws, row, values, alt=(i % 2 == 1))
        ws.cell(row=row - 1, column=3).font = Font(
            bold=True, color=COLOR_EXPENSE if tx["type"] == "expense" else COLOR_INCOME)
        ws.cell(row=row - 1, column=5).number_format = "#,##0.00"
        ws.cell(row=row - 1, column=7).number_format = "#,##0"
    if not rows:
        row = _data_row(ws, row, ["ماكو معاملات بهذي الفترة."] + [""] * (MAIN_COLS - 1))
    return _blank(row)


# ------------------------------------ أقسام 3+4: تصنيفات (مصروف/دخل) ------------------------------------

def _write_category_section(ws, row, title, rows_data, chart_anchor_col="J"):
    row = _section_header(ws, row, title, span=3)
    header_row = row
    row = _table_header(ws, row, ["التصنيف", "عدد العمليات", "الإجمالي"])
    data_start = row
    total_count = total_amount = 0
    for i, c in enumerate(rows_data):
        icon = cat_utils.icon(c["category"])
        row = _data_row(ws, row, [f"{icon} {c['category']}", c["count"], c["total"]], alt=(i % 2 == 1))
        ws.cell(row=row - 1, column=3).number_format = "#,##0"
        total_count += c["count"]
        total_amount += c["total"]
    data_end = row - 1
    row = _data_row(ws, row, ["المجموع", total_count, total_amount], bold_cols=(1, 2, 3),
                     fills={1: COLOR_TOTAL_FILL, 2: COLOR_TOTAL_FILL, 3: COLOR_TOTAL_FILL})
    ws.cell(row=row - 1, column=3).number_format = "#,##0"

    if data_end >= data_start:
        chart = BarChart()
        chart.title = title
        chart.style = 10
        chart.y_axis.title = "المبلغ (IQD)"
        data_ref = Reference(ws, min_col=3, min_row=header_row, max_row=data_end)
        cats_ref = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 15
        chart.height = 8
        ws.add_chart(chart, f"{chart_anchor_col}{header_row}")

    return _blank(row)


# ------------------------------------ قسم 5: الالتزامات ------------------------------------

async def _write_commitments_section(ws, row, user_id):
    row = _section_header(ws, row, "📌 الالتزامات المالية", span=6)
    row = _table_header(ws, row, ["النوع", "الاسم", "الأصلي", "المدفوع", "المتبقي", "الاستقطاع"])
    loans = await db.list_loans(user_id)
    installments = await db.list_installments(user_id)
    i = 0
    for l in loans:
        paid = l["original_amount"] - l["remaining_amount"]
        row = _data_row(ws, row, ["سلفة", l["name"] or f"#{l['id']}", l["original_amount"], paid,
                                   l["remaining_amount"], l["deduction_per_salary"]], alt=(i % 2 == 1))
        for col in (3, 4, 5, 6):
            ws.cell(row=row - 1, column=col).number_format = "#,##0"
        i += 1
    for inst in installments:
        paid = inst["original_amount"] - inst["remaining_amount"]
        row = _data_row(ws, row, ["قسط", inst["name"] or f"#{inst['id']}", inst["original_amount"], paid,
                                   inst["remaining_amount"], inst["installment_amount"]], alt=(i % 2 == 1))
        for col in (3, 4, 5, 6):
            ws.cell(row=row - 1, column=col).number_format = "#,##0"
        i += 1
    if not loans and not installments:
        row = _data_row(ws, row, ["ماكو التزامات مسجّلة."] + [""] * 5)
    return _blank(row)


# ------------------------------------ قسم 6: الدفعات والاستقطاعات ------------------------------------

async def _write_payments_section(ws, row, user_id, d_from, d_to):
    row = _section_header(ws, row, "🔄 الدفعات والاستقطاعات", span=4)
    row = _table_header(ws, row, ["التاريخ", "النوع", "الالتزام", "المبلغ"])
    i = 0
    total = 0.0

    async def _rows_for(items, kind_label, name_fn, list_payments_fn):
        nonlocal row, i, total
        for item in items:
            for p in await list_payments_fn(item["id"], limit=200):
                d = p["created_at"][:10]
                if d_from <= d <= d_to:
                    row = _data_row(ws, row, [d, kind_label, name_fn(item), p["amount"]], alt=(i % 2 == 1))
                    ws.cell(row=row - 1, column=4).number_format = "#,##0"
                    total += p["amount"]
                    i += 1

    await _rows_for(await db.list_loans(user_id), "سلفة", lambda l: l["name"] or f"#{l['id']}", db.list_loan_deductions)
    await _rows_for(await db.list_installments(user_id), "قسط", lambda x: x["name"] or f"#{x['id']}",
                     db.list_installment_payments)
    await _rows_for(await db.list_recurring(user_id), "التزام متكرر", lambda r: r["name"], db.list_recurring_payments)

    if i == 0:
        row = _data_row(ws, row, ["ماكو دفعات مسجّلة بهذي الفترة."] + [""] * 3)
    return _blank(row), total


# ------------------------------------ قسم 7: أهداف الادخار ------------------------------------

async def _write_savings_section(ws, row, user_id, d_from, d_to):
    row = _section_header(ws, row, "🎯 أهداف الادخار", span=5)
    row = _table_header(ws, row, ["الهدف", "المطلوب", "المدخر", "المتبقي", "النسبة"])
    goals = await db.list_savings_goals(user_id)
    total_saved_period = 0.0
    for i, g in enumerate(goals):
        remaining = max(0, g["target_amount"] - g["saved_amount"])
        pct = (g["saved_amount"] / g["target_amount"] * 100) if g["target_amount"] > 0 else 0
        row = _data_row(ws, row, [g["name"], g["target_amount"], g["saved_amount"], remaining, f"{pct:.0f}%"],
                         alt=(i % 2 == 1))
        for col in (2, 3, 4):
            ws.cell(row=row - 1, column=col).number_format = "#,##0"
        for st in await db.list_savings_transactions(g["id"], limit=200):
            d = st["created_at"][:10]
            if d_from <= d <= d_to:
                total_saved_period += st["amount"]
    if not goals:
        row = _data_row(ws, row, ["ماكو أهداف ادخار مسجّلة."] + [""] * 4)
    return _blank(row), total_saved_period


# ------------------------------------ التجميع الكامل (ورقة واحدة) ------------------------------------

async def _build_workbook(user_id, d_from, d_to, title_label):
    rows = await db.transactions_in_range(user_id, d_from, d_to)
    summary = await period_summary(user_id, d_from, d_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "محفظة"
    ws.sheet_view.rightToLeft = True

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=MAIN_COLS)
    title_cell = ws.cell(row=row, column=1, value=f"💰 تقرير المحفظة المالية — {title_label}")
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 30
    row = _blank(row + 1)

    # نحسب مجاميع الاستقطاعات والادخار بالفترة أولًا لقسم الملخص
    total_deductions = await _sum_payments_in_range(user_id, d_from, d_to)
    total_saved_period = await _sum_savings_in_range(user_id, d_from, d_to)

    # الرصيد المرحّل والنهائي — نفس الحسابات المركزية المستخدمة في البوت
    opening = await balance_before(user_id, d_from)
    closing_iqd = opening["net_iqd"] + summary["net_iqd"] - total_deductions

    row = _write_summary_section(ws, row, summary, total_deductions, total_saved_period,
                                 opening_iqd=opening["net_iqd"], closing_iqd=closing_iqd)
    row = _write_transactions_section(ws, row, rows)
    row = _write_category_section(ws, row, "🏷️ التصنيفات (مصروف)", summary["expense_by_category"], "J")
    row = _write_category_section(ws, row, "💰 الدخل حسب التصنيف", summary["income_by_category"], "J")
    row = await _write_commitments_section(ws, row, user_id)
    row, _ = await _write_payments_section(ws, row, user_id, d_from, d_to)
    row, _ = await _write_savings_section(ws, row, user_id, d_from, d_to)

    ws.freeze_panes = "A2"
    widths = [14, 10, 10, 20, 14, 10, 14, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    return wb


async def _sum_payments_in_range(user_id, d_from, d_to):
    total = 0.0
    for l in await db.list_loans(user_id):
        for p in await db.list_loan_deductions(l["id"], limit=200):
            if d_from <= p["created_at"][:10] <= d_to:
                total += p["amount"]
    for inst in await db.list_installments(user_id):
        for p in await db.list_installment_payments(inst["id"], limit=200):
            if d_from <= p["created_at"][:10] <= d_to:
                total += p["amount"]
    for r in await db.list_recurring(user_id):
        for p in await db.list_recurring_payments(r["id"], limit=200):
            if d_from <= p["created_at"][:10] <= d_to:
                total += p["amount"]
    return total


async def _sum_savings_in_range(user_id, d_from, d_to):
    total = 0.0
    for g in await db.list_savings_goals(user_id):
        for st in await db.list_savings_transactions(g["id"], limit=200):
            if d_from <= st["created_at"][:10] <= d_to:
                total += st["amount"]
    return total


async def _export_range(user_id, d_from, d_to, filename_label):
    wb = await _build_workbook(user_id, d_from, d_to, filename_label)
    filename = f"محفظة_{_sanitize_label(filename_label)}.xlsx"
    path = config.EXPORTS_DIR / filename
    wb.save(path)
    return path


async def send_excel_for_user(context: ContextTypes.DEFAULT_TYPE, chat_id: int, user_id: int, label: str):
    """يصدّر كل عمليات مستخدم محدد (كل الأشهر) ويرسل الملف — يُستخدم من لوحة الأدمن."""
    path = await _export_range(user_id, "0000-01-01", "9999-12-31", label)
    await context.bot.send_document(chat_id=chat_id, document=open(path, "rb"), filename=path.name)
    return path


async def handle_excel_callback(update: Update, context: ContextTypes.DEFAULT_TYPE, data: str, user: dict):
    query = update.callback_query
    parts = data.split(":")
    action = parts[1]
    await query.answer("جاري إنشاء الملف...")

    n = dates.now()
    if action == "month":
        months_ago = int(parts[2])
        month = n.month - months_ago
        year = n.year
        while month <= 0:
            month += 12
            year -= 1
        d_from, d_to = dates.month_range(year, month)
        label = f"{ARABIC_MONTHS[month-1]}_{year}"
    else:  # all
        d_from, d_to = "0000-01-01", "9999-12-31"
        label = "كل_الأشهر"

    path = await _export_range(user["id"], d_from, d_to, label)
    await context.bot.send_document(chat_id=update.effective_chat.id, document=open(path, "rb"),
                                     filename=path.name)
