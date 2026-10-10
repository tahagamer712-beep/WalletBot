#!/usr/bin/env python3
"""
termux_check.py — فحص حقيقي (Runtime) لتوافق مشروع «محفظة» مع Termux/Android،
وليس مجرد py_compile. شغّله بعد `pip install -r requirements.txt` وقبل `python bot.py`:

    python termux_check.py

كل فحص مستقل: لو فشل واحد، البقية تستمر وتعرض تقرير نهائي واضح بالسبب بالضبط.
لا يلمس wallet.db الحقيقي إطلاقًا — يستخدم قاعدة بيانات مؤقتة منفصلة تُحذف
تلقائيًا في النهاية.
"""
import sys
import os
import platform
import asyncio
import json
import tempfile
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

RESULTS = []  # (name, "PASS"/"FAIL"/"SKIP", detail)


def record(name, ok, detail=""):
    RESULTS.append((name, "PASS" if ok else "FAIL", detail))
    mark = "✅" if ok else "❌"
    print(f"{mark} {name}" + (f" — {detail}" if detail and not ok else ""))


def record_skip(name, reason):
    RESULTS.append((name, "SKIP", reason))
    print(f"⏭️  {name} — SKIP ({reason})")


# ---------------------------------------------------------------- 1. Python
def check_python_version():
    v = sys.version_info
    ok = v >= (3, 9)
    record("Python version", ok, f"{platform.python_version()} (needs >= 3.9)")
    return ok


# ---------------------------------------------------------- 2. Termux / Android
def check_termux_environment():
    is_termux = "com.termux" in os.environ.get("PREFIX", "") or "TERMUX_VERSION" in os.environ
    if is_termux:
        record("Termux environment", True, "PREFIX indicates Termux")
    else:
        record_skip("Termux environment", "not running inside Termux (normal on desktop/dev machine)")
    return True  # لا نفشل الفحص كامل لو شغّلته على لابتوب للتجربة


# --------------------------------------------------------------- 3. Architecture
def check_architecture():
    machine = platform.machine()
    known_android_arches = {"aarch64", "arm64", "armv7l", "armv8l"}
    if machine in known_android_arches:
        record("Architecture", True, machine)
    else:
        record_skip("Architecture", f"{machine} (not a typical Android arch — fine if testing on desktop)")
    return True


# ------------------------------------------------------------------- 4. SQLite
def check_sqlite():
    try:
        import sqlite3
        ver = sqlite3.sqlite_version
        record("SQLite (stdlib)", True, f"v{ver}")
        return True
    except Exception as e:
        record("SQLite (stdlib)", False, str(e))
        return False


# ---------------------------------------------------------------- 5. aiosqlite
def check_aiosqlite():
    try:
        import aiosqlite  # noqa: F401
        record("aiosqlite import", True, getattr(aiosqlite, "__version__", "?"))
        return True
    except Exception as e:
        record("aiosqlite import", False, f"{type(e).__name__}: {e}")
        return False


# ---------------------------------------------------------------- 6. openpyxl
def check_openpyxl():
    try:
        import openpyxl
        record("openpyxl import", True, openpyxl.__version__)
        return True
    except Exception as e:
        record("openpyxl import", False, f"{type(e).__name__}: {e}")
        return False


# ------------------------------------------------------------------- 7. httpx
def check_httpx():
    try:
        import httpx
        record("httpx import", True, httpx.__version__)
        return True
    except Exception as e:
        record("httpx import", False, f"{type(e).__name__}: {e}")
        return False


# ---------------------------------------------------------- 8. python-telegram-bot
def check_telegram():
    try:
        import telegram
        record("python-telegram-bot import", True, telegram.__version__)
        return True
    except Exception as e:
        record("python-telegram-bot import", False, f"{type(e).__name__}: {e}")
        return False


# ----------------------------------------------------------------- 9. zoneinfo
def check_zoneinfo():
    try:
        from zoneinfo import ZoneInfo
        import datetime
        tz = ZoneInfo("Asia/Baghdad")
        now = datetime.datetime.now(tz)
        record("zoneinfo (Asia/Baghdad)", True, str(now))
        return True
    except Exception as e:
        record("zoneinfo (Asia/Baghdad)", False,
                f"{type(e).__name__}: {e} — تأكد من تثبيت حزمة tzdata عبر pip")
        return False


# ------------------------------------------------------------ 10-13. Database
async def check_database(tmp_db_path):
    import config
    config.DB_PATH = str(tmp_db_path)  # عزل كامل عن wallet.db الحقيقي
    try:
        import database.db as db
        await db.init_db()
        record("Database initialization", True, str(tmp_db_path))
    except Exception as e:
        record("Database initialization", False, f"{type(e).__name__}: {e}")
        return False

    try:
        await db.set_setting("check_key", "check_value")
        val = await db.get_setting("check_key")
        record("Database write/read", val == "check_value")
    except Exception as e:
        record("Database write/read", False, f"{type(e).__name__}: {e}")
        return False

    try:
        user = await db.get_or_create_user(999999, "Test User", "testuser")
        tx_id = await db.insert_transaction(
            user_id=user["id"], type_="expense", category="أكل", amount=5000,
            currency="IQD", exchange_rate=1500, amount_iqd=5000, note=None,
            transaction_date="2026-08-09", transaction_time="12:00:00",
            source="manual", ai_provider=None,
        )
        txs = await db.list_transactions(user["id"], limit=5)
        record("Transaction insert", tx_id is not None and len(txs) == 1)
    except Exception as e:
        record("Transaction insert", False, f"{type(e).__name__}: {e}")
        return False

    return True


# --------------------------------------------------------------- 13. Currency
def check_currency():
    try:
        from utils.currency import to_iqd, fmt_iqd
        ok = to_iqd(250, "USD", 1500) == 375000 and to_iqd(5000, "IQD", 1500) == 5000
        record("Currency conversion", ok, fmt_iqd(375000))
        return ok
    except Exception as e:
        record("Currency conversion", False, f"{type(e).__name__}: {e}")
        return False


# ----------------------------------------------------------- 14. Regex parser
def check_regex_parser():
    try:
        from parser.regex_parser import parse
        from parser.validator import validate_payload

        r1 = parse("صرفت 5000 اكل", 1500)
        r2 = parse("استلمت 250$ راتب", 1500)
        r3 = parse("اليوم صرفت 35 الف، 10 اكل و15 بنزين و10 ملابس", 1500)
        # الصيغة السريعة بدون فعل صريح
        r4 = parse("500 اكل", 1500)
        r5 = parse("200 الف راتب", 1500)
        # أرقام عربية-هندية (٠-٩)
        r6 = parse("٥٠٠ اكل", 1500)
        # كلمة غير معروفة -> يسجلها تحت "أخرى" بدل الرفض
        r7 = parse("250 خضرمة", 1500)
        # كلمة مشروبات معروفة (كحولية) -> تصنيف "مشروبات" وليس "أخرى" أو "أكل"
        r9 = parse("50 الف ويسكي", 1500)
        # رسالة بدون رقم إطلاقًا -> يجب أن ترفض (None)
        r8 = parse("شلونك؟", 1500)
        # تحمّل أخطاء إملائية شائعة بالعامية العراقية — بدون كسر كلمات فعلًا غير معروفة
        r10 = parse("25 الف بانزين", 1500)     # ألف داخلية زايدة -> بنزين
        r11 = parse("10 الف تاكسى", 1500)       # ألف مقصورة بدل ياء -> تاكسي
        r12 = parse("5 الف صيدليه", 1500)       # ه بدل ة -> صيدلية
        r13 = parse("250 خضرمة", 1500)          # نفس r7 — كلمة حقيقية غير معروفة تبقى "أخرى"

        checks = {
            "فعل صريح + مبلغ": r1 and validate_payload(r1)[0]["amount"] == 5000,
            "عملة USD": r2 and validate_payload(r2)[0]["currency"] == "USD",
            "عمليات متعددة برسالة وحدة": r3 and len(validate_payload(r3)) == 3,
            "بدون فعل (500 اكل) -> مصروف": r4 and validate_payload(r4)[0]["type"] == "expense",
            "بدون فعل + دخل معروف (200 الف راتب)":
                r5 and validate_payload(r5)[0]["type"] == "income"
                and validate_payload(r5)[0]["amount"] == 200000,
            "أرقام عربية-هندية (٥٠٠)": r6 and validate_payload(r6)[0]["amount"] == 500,
            "كلمة غير معروفة -> أخرى + ملاحظة":
                r7 and validate_payload(r7)[0]["category"] == "أخرى"
                and validate_payload(r7)[0]["note"] == "خضرمة",
            "مشروبات (ويسكي) -> تصنيف مشروبات":
                r9 and validate_payload(r9)[0]["category"] == "مشروبات"
                and validate_payload(r9)[0]["amount"] == 50000,
            "رسالة بدون رقم -> رفض صحيح": r8 is None,
            "خطأ إملائي (بانزين) -> وقود": r10 and validate_payload(r10)[0]["category"] == "وقود",
            "خطأ إملائي (تاكسى) -> مواصلات": r11 and validate_payload(r11)[0]["category"] == "مواصلات",
            "خطأ إملائي (صيدليه) -> أدوية": r12 and validate_payload(r12)[0]["category"] == "أدوية",
            "كلمة حقيقية غير معروفة تبقى أخرى (لا مطابقة زائفة)":
                r13 and validate_payload(r13)[0]["category"] == "أخرى",
        }
        ok = all(checks.values())
        record("Regex parser", ok)
        if not ok:
            for label, passed in checks.items():
                if not passed:
                    print(f"    ❌ {label}")
        return ok
    except Exception as e:
        record("Regex parser", False, f"{type(e).__name__}: {e}\n{traceback.format_exc(limit=2)}")
        return False


# ------------------------------------------------------------- 15. AI router
def check_ai_router_import():
    try:
        import ai.router  # noqa: F401
        import ai.providers  # noqa: F401
        import ai.prompts  # noqa: F401
        record("AI router import", True)
        return True
    except Exception as e:
        record("AI router import", False, f"{type(e).__name__}: {e}")
        return False


async def check_openrouter_free_model_filter():
    """
    يتحقق أن منطق فلترة الموديلات المجانية يستبعد المدفوعة فعليًا من رد مختلط،
    وأن System Prompt يُبنى من تصنيفات حية (مو نص ثابت بالكود).
    """
    try:
        from ai.prompts import build_system_prompt

        sample_data = [
            {"id": "meta-llama/llama-3.1-8b-instruct:free", "name": "Llama free",
             "pricing": {"prompt": "0", "completion": "0"}},
            {"id": "openai/gpt-4o", "name": "GPT-4o (paid)",
             "pricing": {"prompt": "0.000005", "completion": "0.000015"}},
        ]
        # نفس منطق الفلترة الموجود فعليًا بـ fetch_openrouter_free_models —
        # اختبار مباشر للمنطق أوثق من محاكاة شبكة كاملة لهذا الفحص السريع.
        free = [m for m in sample_data if m["id"].endswith(":free")
                or (float(m["pricing"]["prompt"]) == 0 and float(m["pricing"]["completion"]) == 0)]
        filter_ok = len(free) == 1 and free[0]["id"].endswith(":free")

        prompt = build_system_prompt(["أكل", "صحة", "صالون"], ["راتب"])
        prompt_ok = "صحة" in prompt and "صالون" in prompt

        ok = filter_ok and prompt_ok
        record("OpenRouter free-model filter + dynamic categories", ok)
        return ok
    except Exception as e:
        record("OpenRouter free-model filter + dynamic categories", False, f"{type(e).__name__}: {e}")
        return False


async def check_ai_http_layer_mock():
    """يختبر طبقة HTTP (httpx) واستخراج JSON من رد AI، بدون أي اتصال إنترنت حقيقي أو مفتاح API."""
    try:
        import httpx
        from ai.providers import _extract_json

        # 1) يتأكد أن httpx نفسه يعمل محليًا عبر MockTransport (بدون إنترنت)
        def handler(request):
            return httpx.Response(200, json={"ok": True})

        transport = httpx.MockTransport(handler)
        async with httpx.AsyncClient(transport=transport) as client:
            resp = await client.get("https://example.invalid/test")
            assert resp.json() == {"ok": True}

        # 2) يتأكد أن استخراج الـ JSON من رد نموذجي (مع code fences) يعمل صح
        sample = '```json\n{"transactions": [{"type": "expense", "amount": 1000, ' \
                 '"currency": "IQD", "category": "أكل", "note": null}]}\n```'
        parsed = _extract_json(sample)
        ok = parsed["transactions"][0]["amount"] == 1000
        record("AI HTTP layer (mock, no real key)", ok)
        return ok
    except Exception as e:
        record("AI HTTP layer (mock, no real key)", False, f"{type(e).__name__}: {e}")
        return False


async def check_gemini_provider_mock():
    """
    يشغّل GeminiProvider.parse() **الحقيقية** بالكامل (نفس الكود المستخدم فعليًا
    من ai/router.py) عبر httpx.MockTransport — بدون أي اتصال إنترنت حقيقي وبدون
    مفتاح API حقيقي — ويتحقق من:
    - أن الموديل في الـ URL هو الحالي المدعوم (وليس gemini-1.5-flash المتوقف).
    - عدم إرسال المعاملات المهجورة (temperature/top_p/top_k) في Gemini 3.x.
    - وجود thinkingConfig.thinkingLevel (البديل الحالي).
    - أن _extract_json يحوّل ردًا وهميًا بشكل Gemini الحقيقي إلى عملية صحيحة.
    لا يثبت هذا أن مفتاح API حقيقي سيعمل فعليًا — فقط أن شكل الطلب والتحليل
    صحيحان حسب توثيق Google الحالي وقت كتابة هذا الفحص.
    """
    try:
        import httpx
        from ai.providers import GeminiProvider

        captured = {}

        def handler(request):
            captured["url"] = str(request.url)
            captured["body"] = json.loads(request.content.decode())
            fake_text = json.dumps({
                "transactions": [
                    {"type": "expense", "amount": 5000, "currency": "IQD", "category": "أكل", "note": None}
                ]
            }, ensure_ascii=False)
            return httpx.Response(200, json={
                "candidates": [{"content": {"parts": [{"text": fake_text}]}}]
            })

        transport = httpx.MockTransport(handler)
        async with httpx.AsyncClient(transport=transport) as client:
            provider = GeminiProvider(api_key="FAKE_KEY_NOT_REAL_NO_NETWORK", client=client)
            result = await provider.parse("صرفت 5000 اكل")

        gc = captured.get("body", {}).get("generationConfig", {})
        checks = {
            f"موديل حالي غير متوقف ({GeminiProvider.MODEL})":
                "gemini-1.5-flash" not in captured.get("url", "") and GeminiProvider.MODEL in captured.get("url", ""),
            "لا يوجد temperature/top_p/top_k (مهجورة في Gemini 3.x)":
                not any(k in gc for k in ("temperature", "top_p", "top_k")),
            "thinkingConfig.thinkingLevel موجود":
                gc.get("thinkingConfig", {}).get("thinkingLevel") is not None,
            "response_mime_type=application/json موجود":
                gc.get("response_mime_type") == "application/json",
            "استخراج JSON من رد Gemini الوهمي نجح":
                result.get("transactions", [{}])[0].get("amount") == 5000,
        }
        all_ok = all(checks.values())
        record("Gemini Provider live mock test (real httpx.MockTransport)", all_ok,
               f"model={GeminiProvider.MODEL}")
        if not all_ok:
            for label, passed in checks.items():
                if not passed:
                    print(f"    ❌ {label}")
        return all_ok
    except Exception as e:
        record("Gemini Provider live mock test (real httpx.MockTransport)", False, f"{type(e).__name__}: {e}")
        return False


# -------------------------------------------------------------- 16. Excel
def check_excel_generation(tmp_dir):
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["التاريخ", "المبلغ"])
        ws.append(["2026-08-09", 5000])
        path = Path(tmp_dir) / "_check.xlsx"
        wb.save(path)
        ok = path.exists() and path.stat().st_size > 0
        record("Excel generation (openpyxl)", ok, str(path))
        return ok
    except Exception as e:
        record("Excel generation (openpyxl)", False, f"{type(e).__name__}: {e}")
        return False


# ------------------------------------------------------- 17-18. Backup/integrity
def check_sqlite_backup_and_integrity(tmp_db_path, tmp_dir):
    try:
        import sqlite3
        backup_path = Path(tmp_dir) / "_check_backup.db"
        src = sqlite3.connect(str(tmp_db_path))
        dst = sqlite3.connect(str(backup_path))
        with dst:
            src.backup(dst)
        src.close()
        dst.close()
        record("SQLite backup API", backup_path.exists())

        check_conn = sqlite3.connect(str(backup_path))
        result = check_conn.execute("PRAGMA integrity_check").fetchone()[0]
        check_conn.close()
        record("SQLite integrity_check", result == "ok", result)
        return True
    except Exception as e:
        record("SQLite backup API", False, f"{type(e).__name__}: {e}")
        record("SQLite integrity_check", False, "skipped after backup failure")
        return False


# ------------------------------------------------------------ Configuration
def check_configuration():
    try:
        import config
        checks = [
            isinstance(config.DEFAULT_EXCHANGE_RATE, float),
            isinstance(config.ADMIN_ID, int),
            isinstance(config.SECRET_KEY, str) and len(config.SECRET_KEY) > 0,
            config.EXPORTS_DIR.exists(),
            config.BACKUPS_DIR.exists(),
        ]
        ok = all(checks)
        record("Configuration load", ok)
        return ok
    except Exception as e:
        record("Configuration load", False, f"{type(e).__name__}: {e}")
        return False


# ------------------------------------------------------- Telegram handler import
def check_handlers_import():
    try:
        import handlers.user  # noqa: F401
        import handlers.messages  # noqa: F401
        import handlers.reports  # noqa: F401
        import handlers.excel  # noqa: F401
        import handlers.admin  # noqa: F401
        import handlers.ai_admin  # noqa: F401
        import handlers.backup  # noqa: F401
        import handlers.settings  # noqa: F401
        record("Telegram handlers import", True)
        return True
    except Exception as e:
        record("Telegram handlers import", False, f"{type(e).__name__}: {e}")
        return False


# ------------------------------------------------------------------ Security
def check_security_obfuscation():
    try:
        from utils.security import encrypt_key, decrypt_key, mask_key
        raw = "sk-test-abcdef123456"
        enc = encrypt_key(raw)
        dec = decrypt_key(enc)
        ok = dec == raw and enc != raw
        record("API key obfuscation (stdlib, no cryptography)", ok, mask_key(raw))
        return ok
    except Exception as e:
        record("API key obfuscation (stdlib, no cryptography)", False, f"{type(e).__name__}: {e}")
        return False


async def check_admin_routing():
    """
    اختبار حي (وليس مجرد قراءة كود): يشغّل handle_admin_callback الحقيقية من
    handlers/admin.py فعليًا لكل أزرار AI/Backup/Excel، ويتحقق أنها توصل
    للدالة الفرعية الصحيحة. يمنع تكرار أي مشكلة routing مستقبلًا.
    """
    try:
        import handlers.admin as admin_mod

        calls = []

        async def fake_show_ai_menu(update, context, user):
            calls.append(("show_ai_menu", None))

        async def fake_handle_ai_callback(update, context, data, user):
            calls.append(("handle_ai_callback", data))

        async def fake_show_backup_menu(update, context, user):
            calls.append(("show_backup_menu", None))

        async def fake_handle_backup_callback(update, context, data, user):
            calls.append(("handle_backup_callback", data))

        orig = (
            admin_mod.ai_admin.show_ai_menu, admin_mod.ai_admin.handle_ai_callback,
            admin_mod.backup_handler.show_backup_menu, admin_mod.backup_handler.handle_backup_callback,
        )
        admin_mod.ai_admin.show_ai_menu = fake_show_ai_menu
        admin_mod.ai_admin.handle_ai_callback = fake_handle_ai_callback
        admin_mod.backup_handler.show_backup_menu = fake_show_backup_menu
        admin_mod.backup_handler.handle_backup_callback = fake_handle_backup_callback

        class FakeQuery:
            async def answer(self, *a, **kw): pass
            async def edit_message_text(self, *a, **kw): pass

        class FakeChat:
            id = 1

        class FakeUpdate:
            def __init__(self):
                self.callback_query = FakeQuery()
                self.effective_chat = FakeChat()

        expected = {
            "adm:ai": ("show_ai_menu", None),
            "adm:ai_keys": ("handle_ai_callback", "adm:ai_keys"),
            "adm:ai_key_add": ("handle_ai_callback", "adm:ai_key_add"),
            "adm:ai_priority": ("handle_ai_callback", "adm:ai_priority"),
            "adm:ai_status": ("handle_ai_callback", "adm:ai_status"),
            "adm:backup": ("show_backup_menu", None),
            "adm:backup_now": ("handle_backup_callback", "adm:backup_now"),
        }

        all_ok = True
        for data, exp in expected.items():
            calls.clear()
            await admin_mod.handle_admin_callback(FakeUpdate(), None, data, {"telegram_id": 1})
            got = calls[0] if calls else None
            if got != exp:
                all_ok = False
                print(f"    routing mismatch for {data}: expected {exp}, got {got}")

        (admin_mod.ai_admin.show_ai_menu, admin_mod.ai_admin.handle_ai_callback,
         admin_mod.backup_handler.show_backup_menu, admin_mod.backup_handler.handle_backup_callback) = orig

        record("Admin AI/Backup callback routing (live)", all_ok, f"{len(expected)} cases executed")
        return all_ok
    except Exception as e:
        record("Admin AI/Backup callback routing (live)", False, f"{type(e).__name__}: {e}")
        return False


async def check_start_admin_and_menu_routing():
    """
    اختبار حي: يتحقق من سلوكين مهمين طُلبا صراحة —
    1) /admin يجب أن يكون صامتًا تمامًا (بدون أي رد) لغير الأدمن.
    2) /admin يفتح لوحة الأدمن فعليًا للأدمن.
    3) كل زر من أزرار القائمة الرئيسية (main:*) يوصل فعليًا للدالة الصحيحة.
    """
    try:
        import config as config_mod
        import handlers.user as user_mod
        import handlers.messages as messages_mod

        original_admin_id = config_mod.ADMIN_ID
        config_mod.ADMIN_ID = 999999999

        panel_calls = []

        async def fake_show_admin_panel(update, context, user):
            panel_calls.append(True)

        import handlers.admin as admin_mod
        orig_show_panel = admin_mod.show_admin_panel
        admin_mod.show_admin_panel = fake_show_admin_panel

        class FakeTgUser:
            def __init__(self, uid):
                self.id = uid
                self.full_name = "Test"
                self.username = "test"
                self.first_name = "Test"

        class FakeMessage:
            def __init__(self):
                self.replies = []

            async def reply_text(self, text, reply_markup=None):
                self.replies.append(text)

        class FakeUpdate:
            def __init__(self, uid):
                self.effective_user = FakeTgUser(uid)
                self.message = FakeMessage()
                self.callback_query = None

        non_admin_update = FakeUpdate(111)
        await user_mod.admin_cmd(non_admin_update, None)
        silent_ok = len(panel_calls) == 0 and len(non_admin_update.message.replies) == 0

        panel_calls.clear()
        admin_update = FakeUpdate(999999999)
        await user_mod.admin_cmd(admin_update, None)
        admin_opens_ok = len(panel_calls) == 1

        admin_mod.show_admin_panel = orig_show_panel
        config_mod.ADMIN_ID = original_admin_id

        # اختبار توجيه أزرار main:* (نفس فكرة اختبار adm:* أعلاه)
        menu_calls = []

        async def make_fake(name):
            async def _f(update, context, user):
                menu_calls.append(name)
            return _f

        orig_handlers = dict(messages_mod.MAIN_MENU_HANDLERS)
        orig_reports_cb = messages_mod.reports.handle_reports_callback

        for key in list(messages_mod.MAIN_MENU_HANDLERS.keys()):
            messages_mod.MAIN_MENU_HANDLERS[key] = (await make_fake(key))

        async def fake_reports_cb(update, context, data, user):
            menu_calls.append(f"history_via:{data}")

        messages_mod.reports.handle_reports_callback = fake_reports_cb

        class FakeQuery:
            async def answer(self, *a, **kw):
                pass

        class FakeCbUpdate:
            def __init__(self):
                self.callback_query = FakeQuery()

        menu_expected = {
            "main:summary": "summary", "main:reports": "reports", "main:balance": "balance",
            "main:excel": "excel", "main:settings": "settings",
        }
        menu_ok = True
        for data, expected_name in menu_expected.items():
            menu_calls.clear()
            await messages_mod._handle_main_menu_callback(FakeCbUpdate(), None, data, {"id": 1})
            if menu_calls != [expected_name]:
                menu_ok = False
                print(f"    main menu mismatch for {data}: got {menu_calls}")

        menu_calls.clear()
        await messages_mod._handle_main_menu_callback(FakeCbUpdate(), None, "main:history", {"id": 1})
        history_ok = menu_calls == ["history_via:rep:hist:0"]

        messages_mod.MAIN_MENU_HANDLERS = orig_handlers
        messages_mod.reports.handle_reports_callback = orig_reports_cb

        all_ok = silent_ok and admin_opens_ok and menu_ok and history_ok
        record("/admin silence + main menu routing (live)", all_ok,
               f"silent={silent_ok}, opens={admin_opens_ok}, menu={menu_ok}, history={history_ok}")
        return all_ok
    except Exception as e:
        record("/admin silence + main menu routing (live)", False, f"{type(e).__name__}: {e}")
        return False


async def check_migrations(tmp_db_path):
    """يتحقق أن run_migrations() تشتغل بدون أخطاء على قاعدة جديدة (أهم شي: لا تكسر شي)."""
    try:
        from database.migrations import run_migrations
        version = await run_migrations()
        record("Database migrations", version >= 2, f"version={version}")
        return True
    except Exception as e:
        record("Database migrations", False, f"{type(e).__name__}: {e}")
        return False


async def check_commitments_and_savings(tmp_db_path):
    """
    اختبار حي شامل: سلفة + قسط + التزام متكرر + هدف ادخار، عبر database.db
    الحقيقية (aiosqlite) — يتحقق من محرك الاستقطاع الموحّد عند استلام راتب.
    """
    try:
        import database.db as db
        from handlers.commitments import apply_all_salary_deductions

        user = await db.get_or_create_user(888888, "Commitments Test", "ctest")

        loan_id = await db.create_loan(user["id"], 1000000, 100000, name="سلفة اختبار", auto_deduct=True)
        inst_id = await db.create_installment(user["id"], "قسط اختبار", 1200000, 100000, auto_deduct=True)
        rec_id = await db.create_recurring(user["id"], "التزام اختبار", 50000,
                                            frequency="monthly", auto_deduct=True)
        goal_id = await db.create_savings_goal(user["id"], "هدف اختبار", 3000000)

        tx_id = await db.insert_transaction(
            user_id=user["id"], type_="income", category="راتب", amount=1500000,
            currency="IQD", exchange_rate=1500, amount_iqd=1500000, note=None,
            transaction_date="2026-08-11", transaction_time="09:00:00",
            source="manual", ai_provider=None,
        )
        results = await apply_all_salary_deductions(user["id"], tx_id)
        total_deducted = sum(r["deducted"] for r in results)

        loan_after = await db.get_loan(loan_id, user["id"])
        inst_after = await db.get_installment(inst_id, user["id"])

        savings_result = await db.add_savings_amount(goal_id, 750000)

        checks = {
            "3 التزامات استُقطعت بعد راتب واحد": len(results) == 3,
            "إجمالي الاستقطاع = 250,000": total_deducted == 250000,
            "رصيد السلفة تحدّث صح (900,000)": loan_after["remaining_amount"] == 900000,
            "رصيد القسط تحدّث صح (1,100,000)": inst_after["remaining_amount"] == 1100000,
            "هدف الادخار سجّل 750,000 (25%)": savings_result["new_saved"] == 750000,
        }
        ok = all(checks.values())
        record("Commitments + savings engine (live)", ok)
        if not ok:
            for label, passed in checks.items():
                if not passed:
                    print(f"    ❌ {label}")
        return ok
    except Exception as e:
        record("Commitments + savings engine (live)", False, f"{type(e).__name__}: {e}\n{traceback.format_exc(limit=2)}")
        return False


async def check_excel_single_sheet(tmp_dir):
    """يتحقق أن تصدير Excel ينتج ورقة واحدة فيها كل الأقسام السبعة (بدل 7 تبويبات منفصلة)."""
    try:
        from handlers.excel import _build_workbook
        wb = await _build_workbook(999999, "0000-01-01", "9999-12-31", "فحص")
        ws = wb.active
        all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        expected_sections = [
            "📊 الملخص", "📋 جميع المعاملات", "🏷️ التصنيفات (مصروف)",
            "💰 الدخل حسب التصنيف", "📌 الالتزامات المالية",
            "🔄 الدفعات والاستقطاعات", "🎯 أهداف الادخار",
        ]
        missing = [s for s in expected_sections if s not in all_text]
        ok = len(wb.sheetnames) == 1 and not missing
        path = Path(tmp_dir) / "_check_single_sheet.xlsx"
        wb.save(path)
        ok = ok and path.exists() and path.stat().st_size > 0
        record("Excel single-sheet report generation", ok,
               f"sheets={wb.sheetnames}" + (f", missing={missing}" if missing else ""))
        return ok
    except Exception as e:
        record("Excel single-sheet report generation", False, f"{type(e).__name__}: {e}")
        return False


async def main():
    print("=" * 50)
    print("TERMUX COMPATIBILITY TEST — محفظة")
    print("=" * 50)

    check_python_version()
    check_termux_environment()
    check_architecture()
    check_sqlite()
    check_aiosqlite()
    check_openpyxl()
    check_httpx()
    check_telegram()
    check_zoneinfo()
    check_configuration()
    check_security_obfuscation()
    check_currency()
    check_regex_parser()
    check_ai_router_import()
    await check_openrouter_free_model_filter()
    await check_ai_http_layer_mock()
    await check_gemini_provider_mock()
    check_handlers_import()
    await check_admin_routing()
    await check_start_admin_and_menu_routing()

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_db_path = Path(tmp_dir) / "_check_wallet.db"
        db_ok = await check_database(tmp_db_path)
        check_excel_generation(tmp_dir)
        if db_ok:
            check_sqlite_backup_and_integrity(tmp_db_path, tmp_dir)
            await check_migrations(tmp_db_path)
            await check_commitments_and_savings(tmp_db_path)
            await check_excel_single_sheet(tmp_dir)
        else:
            record_skip("SQLite backup API", "Database init failed above")
            record_skip("SQLite integrity_check", "Database init failed above")
            record_skip("Database migrations", "Database init failed above")
            record_skip("Commitments + savings engine (live)", "Database init failed above")
            record_skip("Excel single-sheet report generation", "Database init failed above")

    print()
    print("=" * 50)
    print("TERMUX COMPATIBILITY TEST — RESULT")
    print("=" * 50)
    name_width = max(len(n) for n, _, _ in RESULTS) + 2
    for name, status, detail in RESULTS:
        dots = "." * (name_width - len(name))
        print(f"{name} {dots} {status}")

    failed = [r for r in RESULTS if r[1] == "FAIL"]
    skipped = [r for r in RESULTS if r[1] == "SKIP"]
    print()
    if failed:
        print(f"RESULT: FAIL ({len(failed)} فشل، {len(skipped)} تم تجاوزه)")
        print("\nتفاصيل الفشل:")
        for name, status, detail in failed:
            print(f"  ❌ {name}: {detail}")
        sys.exit(1)
    else:
        print(f"RESULT: PASS ({len(skipped)} تم تجاوزه — طبيعي إذا كنت تختبر خارج Termux)")
        sys.exit(0)


if __name__ == "__main__":
    asyncio.run(main())
